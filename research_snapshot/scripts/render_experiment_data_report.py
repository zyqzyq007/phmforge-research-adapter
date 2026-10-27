#!/usr/bin/env python3
"""Render the provenance-bound local experiment export as XLSX, HTML and figures."""
import argparse
import csv
import html
import json
import math
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LABELS = {
    'test_joint_correct':'最终联合正确','test_joint_n':'最终联合分母','test_component_correct':'最终原子标注正确','test_component_n':'最终原子标注分母','reference_test_joint_correct':'起点联合正确','reference_test_joint_n':'起点联合分母','reference_test_component_correct':'起点原子标注正确','reference_test_component_n':'起点原子标注分母','id':'记录ID','batch':'批次','dataset':'数据集标识','task':'任务','origin':'数据性质',
    'case':'案例/会话','stage':'阶段','status':'原始运行状态','formal_submission':'正常正式提交',
    'artifact_kind':'制品/成绩口径','unit':'误差单位','scope_limit':'数据与独立性边界',
    'evidence':'证据位置','model':'基础模型/模型','cli':'CLI版本','notes':'限制与判定',
    'test_n':'最终观测数','test_groups':'最终源组/设备数（见边界）','train_n':'训练观测数',
    'dev_n':'开发观测数','train_accuracy':'训练Accuracy','cv_accuracy':'CV Accuracy',
    'dev_accuracy':'开发Accuracy','test_accuracy':'最终Accuracy','reference_test_accuracy':'起始参考Accuracy',
    'train_macro_f1':'训练macro-F1','cv_macro_f1':'CV macro-F1','dev_macro_f1':'开发macro-F1',
    'test_macro_f1':'最终macro-F1','test_balanced_accuracy':'最终Balanced Accuracy',
    'dev_mae':'开发MAE','test_mae':'最终MAE','test_rmse':'最终RMSE','reference_test_mae':'起始参考MAE',
    'test_consistent10_mae':'FEMTO一致10例MAE','test_consistent10_rmse':'FEMTO一致10例RMSE',
    'test_median_error_mm':'定位中位误差mm','test_max_error_mm':'定位最大误差mm',
    'test_cases_within_20mm':'定位20mm内例数','elapsed_seconds':'运行秒数','selected_model':'实际选择',
    'metric_sources':'逐值来源JSON指针','official_benchmark_score':'原版官方成绩',
    'comparison':'对照说明','metric':'指标','before':'保留候选/修改前','after':'提交程序/修改后',
    'delta':'后减前（原单位）','n':'观测数','groups':'源组/设备数','kind':'对照类型',
    'net_extra_errors':'净增加错误窗数','limitation':'解释边界','before_source':'前值来源','after_source':'后值来源',
    'tasks':'涉及任务','phase_rows':'主体阶段数','normal_finished':'正常结束数',
    'formal_submissions':'正式提交数','scored_or_diagnostic_rows':'有数值记录数（含诊断）',
    'extra_infrastructure_starts':'额外基础设施启动','model_requests':'模型请求数','request_scope':'请求计数口径',
    'optimizer_steps':'优化器更新步数','train_groups':'训练源组数','dev_groups':'开发源组数',
    'group_unit':'组的含义','context':'原记录上下文','value':'原始值','absolute_path':'源文件绝对路径',
    'bytes':'源文件字节数','sha256':'SHA256','verification':'核查方式',
    'per_arm_initial_sessions':'每组初选会话','immediate_final_f1':'即时反馈最终F1',
    'delayed_final_f1':'延迟反馈最终F1','immediate_minus_delayed_f1':'即时减延迟F1',
    'interpretation':'解释','limitations':'限制','immediate_source':'即时来源','delayed_source':'延迟来源',
}

MAIN_COLS=['batch','task','case','stage','status','formal_submission','artifact_kind','origin',
    'train_accuracy','cv_accuracy','dev_accuracy','test_accuracy','test_macro_f1','test_balanced_accuracy',
    'test_joint_correct','test_joint_n','test_component_correct','test_component_n','reference_test_joint_correct','reference_test_joint_n','reference_test_component_correct','reference_test_component_n',
    'train_macro_f1','cv_macro_f1','dev_macro_f1','dev_mae','test_mae','test_rmse','unit',
    'test_consistent10_mae','test_consistent10_rmse','test_median_error_mm','test_max_error_mm',
    'test_cases_within_20mm','train_n','dev_n','test_n','test_groups','reference_test_accuracy',
    'reference_test_mae','elapsed_seconds','selected_model','scope_limit','notes','evidence','metric_sources','id']

def val(v):
    if isinstance(v,(dict,list)):return json.dumps(v,ensure_ascii=False)
    return v

def fmt(x,digits=2):
    return '—' if x is None else f'{x:,.{digits}f}'

def score(row,split='test'):
    if row.get(split+'_joint_correct') is not None:
        return f"联合{row[split+'_joint_correct']}/{row[split+'_joint_n']}；原子标注{row[split+'_component_correct']}/{row[split+'_component_n']}"
    if row.get(split+'_accuracy') is not None:
        text=f"{row[split+'_accuracy']*100:.2f}%"
        if row.get(split+'_macro_f1') is not None:text+=f" / F1 {row[split+'_macro_f1']:.4f}"
        return text
    if row.get(split+'_mae') is not None:return f"MAE {row[split+'_mae']:.2f} {row['unit']}"
    if row.get(split+'_median_error_mm') is not None:
        return f"中位 {row[split+'_median_error_mm']:.2f} mm / {row.get(split+'_cases_within_20mm','—')}/7合格"
    return '待评分' if row.get('scoring_status')=='未评分' and split=='test' else '未汇总 / 无适用值'

def evidence_link(ref):
    path=ref.split('#',1)[0]
    if not path.startswith('/'):
        path='/home/jupyter-yyf-zyq/Bo_2/PHM_20261010/'+path
    return path

def make_workbook(data,out):
    wb=Workbook();ws=wb.active;ws.title='阅读说明'
    notes=[
        ('数据时间',data['meta']['created_beijing']),
        ('覆盖范围',data['meta']['scope']),
        ('主体记录',f"{len(data['subjects'])}个实验阶段；含未完成阶段；不等于独立任务/设备数。"),
        ('最新批次',data['meta']['latest_state']),
        ('数值来源','主体与对照数值从已有JSON提取。逐值JSON指针及原文件SHA在后续工作表。'),
        ('结果分母','同一划分的重复会话、窗口、候选不是独立设备；不计算跨任务总准确率或总体失败率。'),
        ('原始值与显示','CSV/JSON中的Accuracy和F1是0–1；Excel的Accuracy列显示百分比，F1保留0–1。'),
        ('RUL单位','FEMTO为秒、XJTU为分钟、C-MAPSS为cycles。不同单位不合并平均。'),
        ('FEMTO分母','主MAE列为全部11例；一致10例成绩单列，不混算。'),
        ('训练与数据真实性','真实拟合不等于实测来源。C-MAPSS是物理仿真；AssetOps fixture现场来源未独立核实。'),
        ('环境与Agent','13条环境/基线结果在单独工作表，不计Agent主体次数；导波数值算法不要求训练神经网络。'),
        ('失败与诊断','预算停止、无完整提交和基础设施故障保留；无提交的参考诊断分不放入主体最终成绩列。'),
        ('历史限制','早期电机/CWRU存在类别总量先验；若仅是旧程序复核，移除元数据不追认独立盲测。'),
        ('科学结论','已保留局部选模/程序退化和反证；“重复反馈导致过拟合”因果与稳定难点尚未确认。'),
        ('本次操作','仅整理已有制品；0新模型调用、0新拟合、0新推理、0新隐藏测试分数。'),
        ('本轮实验状态',data['meta']['latest_state']+'；归档及研究结论另见研究材料。'),
        ('复核方式',f"脚本逐值核对{data['meta']['numeric_cells_verified']}个数字/来源映射；未声称人工审稿或统计显著性。"),
    ]
    ws.append(['项目','说明'])
    for r in notes:ws.append(r)
    ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=115
    for row in ws.iter_rows(min_row=2):
        row[0].font=Font(bold=True,color='17424C');row[1].alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=34
    def table(title,rows,columns=None):
        sheet=wb.create_sheet(title)
        columns=columns or list(dict.fromkeys(k for r in rows for k in r))
        sheet.append([LABELS.get(c,c) for c in columns])
        for r in rows:sheet.append([val(r.get(c)) for c in columns])
        sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions
        for i,c in enumerate(columns,1):
            wide=c in ['notes','scope_limit','evidence','metric_sources','limitation','context','absolute_path','before_source','after_source']
            sheet.column_dimensions[get_column_letter(i)].width=65 if wide else 32 if c in ['case','task','comparison','selected_model','artifact_kind'] else 18
        for row in sheet.iter_rows(min_row=2):
            for col,cell in zip(columns,row):
                cell.alignment=Alignment(vertical='top',wrap_text=col in ['notes','scope_limit','artifact_kind','limitation'])
                if isinstance(cell.value,str):
                    # Force literal text, so exported records cannot become spreadsheet formulas.
                    cell.data_type='s'
                if isinstance(cell.value,float):
                    cell.number_format='0.00%' if col.endswith('accuracy') else '0.0000' if 'f1' in col.lower() else '0.00'
            sheet.row_dimensions[row[0].row].height=42 if title not in ['全部指标','来源索引'] else 20
        sheet.sheet_view.zoomScale=85
        return sheet
    table('批次清单',data['batches'])
    table('逐次Agent结果',data['subjects'],MAIN_COLS)
    table('候选与修改对照',data['comparisons'])
    table('反馈机制对照',data['feedback_controls'])
    table('数据划分',data['splits'])
    table('真实训练资格',data['baselines'],['batch','task','case','origin','test_accuracy','test_macro_f1',
        'test_mae','test_rmse','unit','test_n','optimizer_steps','artifact_kind','notes','evidence','metric_sources'])
    table('全部指标',data['metrics'])
    table('来源索引',data['sources'])
    for sheet in wb:
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.sheet_properties.outlinePr.summaryRight=False
        for cell in sheet[1]:
            cell.fill=PatternFill('solid',fgColor='173D48');cell.font=Font(color='FFFFFF',bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        sheet.row_dimensions[1].height=32
        sheet.sheet_view.showGridLines=False
    path=out/'PHM实验数据总表.xlsx';wb.save(path)
    check=load_workbook(path,read_only=True,data_only=False)
    assert check['逐次Agent结果'].max_row==len(data['subjects'])+1
    assert check['全部指标'].max_row==len(data['metrics'])+1
    count=0
    for expected,actual in zip(data['subjects'],check['逐次Agent结果'].iter_rows(min_row=2,values_only=True)):
        for name,value in zip(MAIN_COLS,actual):
            old=val(expected.get(name))
            if isinstance(old,float):assert math.isclose(old,value,rel_tol=1e-14,abs_tol=1e-14)
            elif old not in (None,''):assert old==value,(name,old,value)
            count+=1
    check.close()
    return count

def make_figure(data,out):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    font=FontProperties(fname='/home/jupyter-yyf-zyq/.fonts/NotoSansCJKsc-Regular.otf')
    plt.rcParams['font.family']=font.get_name();plt.rcParams['axes.unicode_minus']=False
    rows=[next(r for r in data['comparisons'] if r['id']==id_) for id_ in ['A305_a_review','A307_pdm_cost_benefit_002','A314_rotor_component','A314_pdm_safety_004']]
    r=next(r for r in data['comparisons'] if r['id']=='A315_xjtu_candidate')
    fig,ax=plt.subplots(1,2,figsize=(13.2,4.7),gridspec_kw={'width_ratios':[2.2,1]},facecolor='#FAFCFD')
    colors=['#176B70','#DD795F']
    for a in ax:
        a.set_facecolor('#FAFCFD');a.spines[['top','right']].set_visible(False)
        a.grid(axis='x',alpha=.15);a.set_axisbelow(True)
    for i,v in enumerate(rows):
        for offset,key,color in [(-.16,'before',colors[0]),(.16,'after',colors[1])]:
            value=v[key]*100;ax[0].barh(i+offset,value,height=.28,color=color)
            ax[0].text(value+1,i+offset,f'{value:.2f}%',va='center',fontsize=10)
    ax[0].set_yticks(range(4),['转子 · A305复核','CWRU · A307','转子组件 · A314','转子起始参考 · A314'],fontproperties=font)
    ax[0].invert_yaxis();ax[0].set_xlim(0,114);ax[0].set_xlabel('隐藏测试 Accuracy（越高越好）',fontproperties=font)
    ax[0].set_title('已有有效候选与提交程序',loc='left',fontproperties=font,pad=17)
    for i,(key,color) in enumerate(zip(['before','after'],colors)):
        ax[1].barh(i,r[key],height=.42,color=color);ax[1].text(r[key]+8,i,f'{r[key]:.2f}',va='center',fontsize=10)
    ax[1].set_yticks([0,1],['已检查的 KNN','新拟合 ET'],fontproperties=font);ax[1].invert_yaxis()
    ax[1].set_xlim(0,540);ax[1].set_xlabel('XJTU 按轴承平均 MAE / 分钟\n（越低越好）',fontproperties=font)
    ax[1].set_title('寿命选模损失 · A315',loc='left',fontproperties=font,pad=17)
    from matplotlib.patches import Patch
    fig.legend(handles=[Patch(color=colors[0],label='可用/实际检查过的候选'),Patch(color=colors[1],label='Agent 提交程序')],
               loc='lower center',ncol=2,frameon=False,prop=font,bbox_to_anchor=(.5,.05))
    fig.text(.035,.015,'5 条已选局部轨迹，来自4批实验；不是随机代表样本，不能估计总体失败率或证明统一机制。',fontproperties=font,fontsize=9,color='#52636B')
    fig.tight_layout(rect=[0,.13,1,1]);fig.savefig(out/'selection_comparisons.png',dpi=180,bbox_inches='tight')
    fig.savefig(out/'selection_comparisons.svg',bbox_inches='tight');plt.close(fig)

HTML = r'''<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>PHM 真实实验数据总表</title>
<style>
:root{--ink:#18303b;--muted:#63727d;--line:#dfe7eb;--teal:#176b70;--pale:#edf6f5;--orange:#ad5035}*{box-sizing:border-box}body{margin:0;background:#f4f7f9;color:var(--ink);font:15px/1.6 system-ui,"Noto Sans CJK SC","Microsoft YaHei",sans-serif}a{color:var(--teal)}main{max-width:1450px;margin:auto;padding:32px 34px 64px}.eyebrow{font-size:12px;letter-spacing:1.8px;color:var(--teal);font-weight:700}h1{font-size:30px;letter-spacing:-.5px;margin:7px 0}h2{font-size:19px;margin:0 0 12px}.sub{color:var(--muted);max-width:1040px;margin:8px 0 18px}.links{display:flex;gap:10px;flex-wrap:wrap}.btn{display:inline-block;text-decoration:none;border:1px solid var(--line);border-radius:8px;padding:9px 14px;background:white;font-weight:600}.btn.primary{background:var(--teal);color:white;border-color:var(--teal)}.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:24px 0}.kpi,.card{background:white;border:1px solid var(--line);border-radius:12px;padding:18px}.kpi strong{font-size:29px;line-height:1.2;display:block;margin:5px 0}.kpi span,.muted{color:var(--muted);font-size:13px}.notice{border-left:4px solid var(--teal);padding:12px 16px;background:var(--pale);border-radius:5px;font-size:14px}.tabs{display:flex;gap:7px;flex-wrap:wrap;margin:24px 0 15px}.tabs button{border:0;background:#e8eef1;color:var(--ink);border-radius:8px;padding:10px 16px;cursor:pointer;font:inherit}.tabs button.active{background:var(--ink);color:white}.panel{display:none}.panel.active{display:block}.figure{width:100%;height:auto;display:block}.grid{display:grid;grid-template-columns:1fr 1fr;gap:15px;margin:15px 0}.table-wrap{overflow:auto;border:1px solid var(--line);border-radius:9px;background:white}table{border-collapse:collapse;width:100%;font-size:13px}th{background:#edf2f4;text-align:left;white-space:nowrap;font-weight:650}td,th{padding:11px 13px;border-bottom:1px solid var(--line);vertical-align:top}tbody tr:hover{background:#f9fbfc}td.num{font-variant-numeric:tabular-nums;white-space:nowrap}.filters{display:flex;gap:9px;flex-wrap:wrap;align-items:center;margin:12px 0}.filters input,.filters select{background:white;border:1px solid #ccd8dd;border-radius:7px;padding:9px 11px;font:inherit;font-size:13px}.filters input{flex:1;min-width:200px}.tag{display:inline-block;border-radius:4px;background:#eaf3f1;color:#215b53;padding:2px 6px;font-size:11px;white-space:nowrap}.tag.warn{background:#fbede6;color:var(--orange)}.tag.pending{background:#eaf0fa;color:#385c9c}details summary{cursor:pointer;color:var(--teal)}.details{min-width:320px;max-width:600px;font-size:12px;word-break:break-word}.score{font-weight:650;white-space:nowrap}.legend{font-size:12px;color:var(--muted);margin:8px 0}.source{font-family:ui-monospace,monospace;font-size:11px;overflow-wrap:anywhere}.foot{margin-top:20px;color:var(--muted);font-size:12px}.count{font-size:13px;color:var(--muted)}@media(max-width:800px){main{padding:20px 14px}.kpis{grid-template-columns:1fr 1fr}.grid{grid-template-columns:1fr}h1{font-size:25px}}@media print{.tabs,.filters,.links{display:none}.panel{display:block;break-before:page}.table-wrap{overflow:visible}main{padding:0}.kpis{grid-template-columns:repeat(4,1fr)}.source{font-size:8px}}
</style></head><body><main>
<div class="eyebrow">PHM RESEARCH · EVIDENCE SNAPSHOT</div><h1>真实实验数据总表</h1><p class="sub" id="snapshot"></p>
<div class="links"><button class="btn" onclick="document.getElementById('benchDialog').showModal()">本轮Bench与真实流程</button><a class="btn primary" href="PHM实验数据总表.xlsx" download>下载 Excel 总表</a><a class="btn" href="agent_results.csv" download>逐次结果 CSV</a><a class="btn" href="experiment_data.json" download>完整结构化 JSON</a><a class="btn" href="metric_details.csv" download>已提取标量指标 CSV</a></div>
<div class="kpis" id="kpis"></div><div class="notice">已导入后续真实外评，<b>成功、退化和未完成阶段全部保留</b>。仍未确认“重复反馈导致过拟合”的因果或困难机制。可执行实验均已结束；55个数值方向、51正常外评，研究适配边界见下方台账。</div>
<dialog id="benchDialog" style="border:1px solid #dfe7eb;border-radius:12px;max-width:920px;width:90%;padding:26px;color:#18303b"><button class="btn" style="float:right" onclick="this.closest('dialog').close()">关闭</button><h2>本轮PHMForge：输入、工具与评价</h2><p>冻结公开75题。已尝试55个具备条件的数值方向：23个实测／实测派生、4个原始NASA FD001、28个作者N-CMAPSS派生QA。其余16个原目标缺件、4个明确合成题保留在完整台账。</p><p>任务原文＋允许的训练／开发资料 → Agent反复推理、取数、调用MCP／Python、真实拟合与验证 → 封存程序 → 隐藏查询真实推理 → 外部固定真值数值评价。</p><div class="table-wrap"><table><thead><tr><th>任务方向</th><th>输入</th><th>输出与评价</th></tr></thead><tbody><tr><td>分类13题</td><td>真实振动／电流源或统计特征、允许的标签</td><td>逐ID类别；accuracy、F1和完整记录指标。12正常、1预算终态。</td></tr><tr><td>RUL3题</td><td>FD001／FEMTO／XJTU的训练历史与当前前缀</td><td>逐单位RUL；MAE、RMSE。cycles／秒／分钟分别统计。</td></tr><tr><td>成本／安全11题</td><td>上述预测任务与原题已给参数</td><td>可核查预测及费用算术；未提供风险／合规／动作效果时保留未知。</td></tr><tr><td>Engine28题</td><td>600×33源记录，或10个60×33前缀的有序元组</td><td>5组件或状态／区间的源QA标注一致性；25正常、3预算终态，不是物理概率校准。</td></tr></tbody></table></div><p>主体使用deepseek-flash与原生coding CLI；可用文件、Python及22个MCP入口。缺有效数据合同的工具明确不可用，不以常数或伪训练补足。真正需要训练时实际fit／更新参数；允许复用已验证模型。隐藏最终标签不提供给开发Agent。</p><p>本轮研究的是coding Agent开发与选择PHM程序，不等于在线机务Agent已实施维护，也不称完整原版官方成绩。185阶段包含历史重复及4条导波；完整来源与限制可在各表展开。</p></dialog><nav class="tabs"><button class="active" data-tab="overview">任务概览</button><button data-tab="runs">逐次实验</button><button data-tab="contrasts">候选 / 修改对照</button><button data-tab="baseline">拟合凭证样例</button><button data-tab="sources">口径与来源</button></nav>
<section id="overview" class="panel active"><div class="card" style="margin-top:15px"><h2>PHMForge全部75题</h2><p id="coverageStatus" class="legend"></p><details><summary>展开逐题状态与缺口</summary><div class="table-wrap"><table id="queueTable"></table></div></details></div><div class="card"><h2>局部程序变化：退化、持平与改善</h2><img class="figure" src="selection_comparisons.png" alt="四条分类轨迹及一条寿命轨迹的可用候选与最终提交对照"><p class="legend">这5条轨迹不是随机代表样本；重复会话和相关窗口不计独立设备。成功、阴性结果及未完成阶段均保留。转子组件对照按整记录多数票均20/20，窗口差异不等于设备级损害。XJTU图只比较MAE；高估与低估误差存在取舍，不等于安全效用。</p></div><div class="card" style="margin-top:15px"><h2>已评分的任务族</h2><div class="table-wrap"><table id="taskTable"></table></div></div><div class="card" style="margin-top:15px"><h2>批次与失败尝试清单</h2><p class="legend">主体阶段数含未完成阶段；额外基础设施启动单列。模型请求数缺失表示本表未聚合，不表示0。</p><div class="table-wrap"><table id="batchTable"></table></div></div></section>
<section id="runs" class="panel"><div class="card"><h2>逐次 Agent 结果</h2><p class="legend">Accuracy 显示百分比，F1为0–1；RUL保留各自单位。空值表示未汇总、不适用或未评分。展开“证据”查看逐值来源。</p><div class="filters"><input id="search" placeholder="搜索任务、批次、案例或限制…"><select id="dataset"><option value="">全部任务</option></select><select id="batch"><option value="">全部批次</option></select><select id="status"><option value="">全部状态</option><option value="finished">正常结束</option><option value="other">预算 / 超时 / 未完成</option><option value="pending">隐藏测试待评分</option></select><button class="btn" id="reset">重置</button></div><p class="count" id="runCount"></p><div class="table-wrap"><table id="runTable"></table></div></div></section>
<section id="contrasts" class="panel"><div class="card"><h2>候选与程序修改对照</h2><p class="legend">Accuracy 越高越好；MAE越低越好。“后减前”使用原单位。组件消融、实际候选和匹配重建不能混成同一种证据。</p><div class="table-wrap"><table id="contrastTable"></table></div></div></section>
<section id="baseline" class="panel"><div class="card"><h2>真实拟合与环境资格样例</h2><p class="legend">这13条既有资格样例不是全部拟合次数，也不是Agent主体；新增参考及权重凭证见各批归档。真实训练不等于实测数据：C-MAPSS为物理仿真；AssetOps冷水机fixture现场来源尚未核实。</p><div class="table-wrap"><table id="baselineTable"></table></div></div><div class="card" style="margin-top:15px"><h2>已直接提取的划分计数</h2><p class="legend">这里的“组”遵循各批原定义；缺项见逐次记录及协议，不补造设备身份。</p><div class="table-wrap"><table id="splitTable"></table></div></div></section>
<section id="sources" class="panel"><div class="grid"><div class="card"><h2>阅读口径</h2><ul><li>覆盖A227导波及后续已评分真实计算研究线，未混入此前333题原工具批量分数。</li><li>研究适配成绩不称原版PHMForge官方成绩。</li><li>FEMTO全部11例与一致10例分别保留在Excel。</li><li>XJTU主MAE按轴承平均；15时点只有3台轴承。</li><li>早期类别总量先验、无完整提交、MCP限制均在逐次记录备注中。</li><li>禁止把窗口/会话/候选相加成独立设备或总体失败率。</li><li>本次导出0新训练／模型调用／推理；包含此前已经真实完成的外评。</li></ul></div><div class="card"><h2>可复核文件</h2><p>逐次结果、训练资格、候选对照、批次清单、数据划分、全部指标与来源索引均提供CSV；Excel保留筛选、冻结表头和原始精度。</p><p><a href="comparisons.csv">候选对照CSV</a> · <a href="batch_inventory.csv">批次CSV</a> · <a href="data_splits.csv">划分CSV</a> · <a href="baseline_results.csv">资格CSV</a> · <a href="evidence_sources.csv">来源索引</a></p><p><a href="selection_comparisons.svg">可导出SVG图</a> · <a href="selection_comparisons.png">PNG图</a> · <a href="validation.json">导出核验回执</a></p><p class="muted">来源链接指向当前研究机器的真实文件；跨机器使用时请按索引取得对应归档。SHA用于检查源文件身份。</p></div></div><div class="table-wrap"><table id="sourceTable"></table></div></section>
<p class="foot">本轮可执行数值实验已结束；研究材料与验收回执保留。研究适配不等于完整官方复现，不新增因果结论、显著性或总体发生率。</p>
</main><script id="data" type="application/json">__DATA__</script><script>
const D=JSON.parse(document.querySelector('#data').textContent),E=s=>String(s??'—').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const num=(v,n=2)=>v===null||v===undefined?'—':Number(v).toFixed(n),pct=v=>v==null?'—':num(v*100)+'%';
function score(r,p='test'){if(r[p+'_joint_correct']!=null)return '联合 '+r[p+'_joint_correct']+'/'+r[p+'_joint_n']+' · 标注 '+r[p+'_component_correct']+'/'+r[p+'_component_n'];if(r[p+'_accuracy']!=null)return pct(r[p+'_accuracy'])+(r[p+'_macro_f1']!=null?' · F1 '+num(r[p+'_macro_f1'],4):'');if(r[p+'_mae']!=null)return 'MAE '+num(r[p+'_mae'])+' '+E(r.unit);if(r[p+'_median_error_mm']!=null)return '中位 '+num(r[p+'_median_error_mm'])+' mm · '+r[p+'_cases_within_20mm']+'/7';return p==='test'&&r.scoring_status==='未评分'?'待评分':'—'}
function table(id,heads,rows){document.getElementById(id).innerHTML='<thead><tr>'+heads.map(x=>'<th>'+x+'</th>').join('')+'</tr></thead><tbody>'+rows.map(row=>'<tr>'+row.map(x=>'<td>'+x+'</td>').join('')+'</tr>').join('')+'</tbody>'}
function link(ref,label='源文件'){const p=ref.split('#')[0],a=p.startsWith('/')?p:'/home/jupyter-yyf-zyq/Bo_2/PHM_20261010/'+p;return '<a href="'+E(a)+'">'+E(label)+'</a>'}
const normal=D.subjects.filter(r=>r.status==='finished').length,available=D.subjects.filter(r=>r.test_accuracy!=null||r.test_mae!=null||r.test_median_error_mm!=null||r.test_joint_correct!=null).length;
document.querySelector('#snapshot').textContent='数据快照 '+D.meta.created_beijing.replace('T',' ').slice(0,19)+'（北京时间） · '+D.subjects.length+'个历史主体阶段，含重复／补全和未完成；不是独立任务或设备数。';
document.querySelector('#kpis').innerHTML=[[D.meta.phmforge_attempted_directions+'/75','已真实尝试的方向','全部具备数值执行条件的方向'],[D.meta.phmforge_normally_scored_directions,'正常交付外评','交付不等于预测正确'],[D.meta.budget_terminal_directions,'预算未交付方向','保留终态，不给起始程序冒记分'],[D.meta.qualified_gaps+' + '+D.meta.explicit_synthetic_excluded,'原目标缺口＋明确合成题','完整75题逐项保留']].map(v=>'<div class="kpi"><span>'+v[1]+'</span><strong>'+v[0]+'</strong><span>'+v[2]+'</span></div>').join('');
document.querySelector('#coverageStatus').textContent=D.meta.phmforge_attempted_directions+'/75个原题方向已开展真实实验，'+D.meta.phmforge_normally_scored_directions+'个有正常交付外评。可执行数值清单'+D.meta.potential_numerical_directions+'个，包含研究适配及部分辅助合同；其余缺口与明确合成题逐项保留。';
table('queueTable',['原题','题面数据 / 配置数据','当前状态','支持范围与缺口'],D.phmforge_queue.rows.map(r=>[E(r.original_task_id),E(r.prompt_dataset)+' / '+E(r.configured_dataset),E(r.current_state),E(r.reason)+(r.evidence?'<br>'+link(r.evidence):'')]));
const datasets=[...new Set(D.subjects.map(r=>r.dataset))];
table('taskTable',['任务','数据性质','阶段记录','批次','数据与独立性边界'],datasets.map(k=>{const rs=D.subjects.filter(r=>r.dataset===k),t=D.tasks[k]||[k,'见原记录','见原记录'];return [E(t[0]),E(t[1]),rs.length,[...new Set(rs.map(r=>r.batch))].join('、'),E(t[2])]}));
table('batchTable',['批次','主体阶段','正常结束','额外基础设施启动','模型请求','进度与解释'],D.batches.map(r=>[E(r.batch),r.phase_rows,r.normal_finished,r.extra_infrastructure_starts,r.model_requests??'未汇总',E(r.notes)+'<br>'+link(r.evidence,'批次报告')]));
for(const k of datasets){let o=new Option((D.tasks[k]||[k])[0],k);document.querySelector('#dataset').add(o)}
for(const k of [...new Set(D.subjects.map(r=>r.batch))])document.querySelector('#batch').add(new Option(k,k));
function renderRuns(){const q=document.querySelector('#search').value.toLowerCase(),ds=document.querySelector('#dataset').value,b=document.querySelector('#batch').value,st=document.querySelector('#status').value;const rows=D.subjects.filter(r=>(!q||JSON.stringify([r.task,r.batch,r.case,r.notes,r.scope_limit]).toLowerCase().includes(q))&&(!ds||r.dataset===ds)&&(!b||r.batch===b)&&(!st||(st==='pending'?r.scoring_status==='未评分':st==='other'?r.status!=='finished':r.status==='finished')));document.querySelector('#runCount').textContent='显示 '+rows.length+' / '+D.subjects.length+' 条阶段记录';table('runTable',['任务 / 批次','案例 / 阶段','运行与成绩口径','开发','最终结果','最终观测 / 组','证据与限制'],rows.map(r=>[E(r.task)+'<br><span class="muted">'+E(r.batch)+' · '+E(r.origin)+'</span>',E(r.case)+'<br>'+E(r.stage),'<span class="tag '+(r.scoring_status==='未评分'?'pending':r.status!=='finished'?'warn':'')+'">'+E(r.status)+'</span><br><span class="muted">'+E(r.artifact_kind)+'</span>','<span class="score">'+score(r,'dev')+'</span>','<span class="score">'+score(r)+'</span>',E(r.test_n)+' / '+E(r.test_groups),'<details><summary>查看证据</summary><div class="details"><p>'+E(r.notes)+'</p><p>'+E(r.scope_limit)+'</p><p>训练：'+score(r,'train')+'；CV F1：'+num(r.cv_macro_f1,4)+'</p><p>'+link(r.evidence)+'</p><div class="source">'+E(r.evidence)+'</div><p class="source">'+E(JSON.stringify(r.metric_sources))+'</p></div></details>']))}
for(const id of ['search','dataset','batch','status'])document.getElementById(id).addEventListener('input',renderRuns);document.querySelector('#reset').onclick=()=>{for(const id of ['search','dataset','batch','status'])document.getElementById(id).value='';renderRuns()};renderRuns();
table('contrastTable',['任务 / 批次','对照','前 → 后','后减前','观测 / 组','证据类型与边界'],D.comparisons.map(r=>[E(r.task)+'<br>'+E(r.batch),E(r.comparison),r.unit==='比例'?pct(r.before)+' → '+pct(r.after):num(r.before)+' → '+num(r.after)+' '+E(r.unit),r.unit==='比例'?num(r.delta*100)+' 个百分点':num(r.delta)+' '+E(r.unit),E(r.n)+' / '+E(r.groups),E(r.kind)+'<br><span class="muted">'+E(r.limitation)+'</span><br>'+link(r.before_source)]));
table('baselineTable',['任务 / 模型','数据性质','实际结果','优化更新','解释'],D.baselines.map(r=>[E(r.task)+'<br>'+E(r.case),E(r.origin),score(r),E(r.optimizer_steps),E(r.notes)+'<br>'+link(r.evidence)]));
table('splitTable',['批次 / 划分','训练 / 开发 / 最终观测','训练 / 开发 / 最终组','组的含义'],D.splits.map(r=>[E(r.batch)+' / '+E(r.case),[r.train_n,r.dev_n,r.test_n].map(E).join(' / '),[r.train_groups,r.dev_groups,r.test_groups].map(E).join(' / '),E(r.group_unit)+'<br>'+link(r.evidence)]));
table('sourceTable',['源文件','字节数','SHA256'],D.sources.map(r=>[link(r.absolute_path,r.id),r.bytes,'<span class="source">'+E(r.sha256)+'</span>']));
document.querySelectorAll('.tabs button').forEach(b=>b.onclick=()=>{document.querySelectorAll('.tabs button,.panel').forEach(e=>e.classList.remove('active'));b.classList.add('active');document.getElementById(b.dataset.tab).classList.add('active')});
</script></body></html>'''

def main():
    p=argparse.ArgumentParser();p.add_argument('--dir',type=Path,required=True);args=p.parse_args();out=args.dir
    data=json.loads((out/'experiment_data.json').read_text())
    verified=make_workbook(data,out);make_figure(data,out)
    embedded=json.dumps(data,ensure_ascii=False,separators=(',',':')).replace('</','<\\/')
    page=HTML.replace('__DATA__',embedded)
    feedback_card='<div class="card" style="margin-top:15px"><h2>反证：反馈时机对照</h2><p class="legend">每批即时/延迟各3次初选。当前零差值不等于证明不存在效应；各批范围限制保留。</p><div class="table-wrap"><table><thead><tr><th>任务</th><th>即时反馈最终F1</th><th>延迟反馈最终F1</th><th>差值</th><th>限制</th></tr></thead><tbody>'
    for r in data['feedback_controls']:
        feedback_card+='<tr><td>'+html.escape(r['task']+' / '+r['batch'])+'</td><td>'+fmt(r['immediate_final_f1'],4)+'</td><td>'+fmt(r['delayed_final_f1'],4)+'</td><td>'+fmt(r['immediate_minus_delayed_f1'],4)+'</td><td>'+html.escape(r['limitations'])+'</td></tr>'
    feedback_card+='</tbody></table></div></div>'
    feedback_card+='<p><a href="research_assessment.md">研究结论、反证与下一步验证计划</a> · <a href="phmforge75_real_coverage.csv">75题完整台账CSV</a></p>'
    anchor='<div class="card" style="margin-top:15px"><h2>批次与失败尝试清单'
    assert page.count(anchor)==1
    page=page.replace(anchor,feedback_card+anchor)
    (out/'index.html').write_text(page)
    # Verify the HTML carries exactly the exported machine-readable records.
    from bs4 import BeautifulSoup
    soup=BeautifulSoup((out/'index.html').read_text(),'html.parser')
    embedded_data=json.loads(soup.select_one('#data').string)
    assert embedded_data['subjects']==data['subjects']
    assert soup.find('h2',string='反证：反馈时机对照') is not None
    assert soup.select_one('a[href="research_assessment.md"]') is not None
    refs=[]
    for a in soup.find_all(['a','img']):
        ref=a.get('href',a.get('src',''))
        if ref and not ref.startswith(('/','http','#')):
            refs.append(ref)
    missing=[r for r in refs if r!='validation.json' and not (out/r).exists()]
    assert not missing,missing
    validation=dict(status='passed',source_value_checks=data['meta']['numeric_cells_verified'],
        excel_cell_roundtrip_checks=verified,html_json_matches=True,local_download_links_checked=len(refs),
        subject_rows=len(data['subjects']),baseline_rows=len(data['baselines']),comparison_rows=len(data['comparisons']),
        metrics_rows=len(data['metrics']),new_fits=0,new_LLM_calls=0,new_hidden_scores=0,
        limitations=['HTML数值和链接静态检查；浏览器交互需另行检查','非论文人工审稿签署'])
    (out/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(validation,ensure_ascii=False))

if __name__=='__main__':main()
